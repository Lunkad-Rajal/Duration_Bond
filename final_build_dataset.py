import datetime as dt
import logging
import re
import warnings
from pathlib import Path

import numpy as np
import pandas as pd

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger("final_dataset")

RAW_DIR = Path(r"C:\A - Personal\User\UnderGrad\MBFM Project\Datasets\raw_data")
OUT_DIR = Path(r"C:\A - Personal\User\UnderGrad\MBFM Project\Datasets\processed")
RBI_GSEC_XLSX = RAW_DIR / "Secondary Market Outright Transactions in Government Securities (Face Value).xlsx"

CBM_COLUMN_MAP = {
    "Trade Date": "date",
    "ISIN": "isin",
    "Last Trade Price (in Rs.)": "last_price",
    "Last Trade Value (Rs. in lacs)": "last_trade_value",
    "Total Trade Value (Rs. in lacs)": "traded_value",
    "Last Trade Yield (YTM) (Annualized) (%)": "last_yield",
    "Weighted Average Price  (Rs.)": "vwap_price",
    "Weighted Average Yield (YTM) (%)": "vwap_yield",
}
CBM_REQUIRED = set(CBM_COLUMN_MAP.values())

YIELD_MIN, YIELD_MAX = 0.5, 20.0
PRICE_MIN, PRICE_MAX = 40.0, 180.0
MAX_ABS_DY = 5.0
GSEC_MAX_FILL_DAYS = 10
GSEC_MIN_DATE = dt.datetime(2017, 1, 1)
HORIZONS = (5, 10, 20, 40, 60)
STALE_GAP = 20
EXCLUDED_SEC_TYPES = {"CP", "CD"}
FREQ_MAP = {"yearly": 1, "annual": 1, "halfyearly": 2, "semiannual": 2, "quarterly": 4, "monthly": 12,
            "1": 1, "2": 2, "4": 4, "12": 12}


def read_cbm_file(path: Path):
    try:
        df = pd.read_csv(path)
    except Exception as e:
        log.warning("Could not read %s: %s", path.name, e)
        return None
    df = df.rename(columns=CBM_COLUMN_MAP)
    missing = CBM_REQUIRED - set(df.columns)
    if missing:
        log.warning("%s missing %s. Columns found: %s", path.name, sorted(missing), list(df.columns))
        return None
    return df[list(CBM_REQUIRED)]


def load_cbm_panel(raw_dir: Path) -> pd.DataFrame:
    files = sorted(raw_dir.rglob("CBM_*.csv"))
    log.info("Loading %d CBM daily files", len(files))
    frames = [f for f in (read_cbm_file(p) for p in files) if f is not None]
    if not frames:
        raise SystemExit(f"No usable CBM_*.csv files found under {raw_dir}")
    panel = pd.concat(frames, ignore_index=True)
    panel["isin"] = panel["isin"].astype(str).str.strip()
    panel["date"] = pd.to_datetime(panel["date"], format="%d-%b-%Y", errors="coerce").astype("datetime64[ns]")
    for col in ["last_price", "last_trade_value", "traded_value", "last_yield", "vwap_price", "vwap_yield"]:
        panel[col] = pd.to_numeric(panel[col], errors="coerce")
    log.info("Raw CBM rows: %d", len(panel))
    return panel


def clean_panel(panel: pd.DataFrame) -> pd.DataFrame:
    before = len(panel)
    panel = panel.dropna(subset=["date", "vwap_yield", "vwap_price", "traded_value"])
    panel = panel[
        panel["vwap_yield"].between(YIELD_MIN, YIELD_MAX)
        & panel["vwap_price"].between(PRICE_MIN, PRICE_MAX)
        & (panel["traded_value"] > 0)
        & panel["isin"].str.startswith("INE")
    ]
    panel = panel.sort_values(["isin", "date"]).drop_duplicates(["isin", "date"], keep="last")
    log.info("Filters: %d -> %d rows, %d ISINs", before, len(panel), panel["isin"].nunique())
    return panel.reset_index(drop=True)


def drop_excluded_types(panel: pd.DataFrame, master) -> pd.DataFrame:
    if master is None or "sec_type" not in master.columns:
        return panel
    excluded = set(master.loc[master["sec_type"].isin(EXCLUDED_SEC_TYPES), "isin"])
    out = panel[~panel["isin"].isin(excluded)].reset_index(drop=True)
    log.info("Dropped commercial paper / CDs: %d rows, %d ISINs", len(panel) - len(out),
             panel.loc[panel["isin"].isin(excluded), "isin"].nunique())
    return out


def build_events(panel: pd.DataFrame) -> pd.DataFrame:
    grp = panel.groupby("isin")
    events = panel.copy()
    events["prev_date"] = grp["date"].shift(1)
    events["prev_yield"] = grp["vwap_yield"].shift(1)
    events["prev_price"] = grp["vwap_price"].shift(1)
    events = events.dropna(subset=["prev_date"])
    events["gap"] = (events["date"] - events["prev_date"]).dt.days
    events["raw_dy"] = events["vwap_yield"] - events["prev_yield"]
    events = events[events["raw_dy"].abs() < MAX_ABS_DY].reset_index(drop=True)
    log.info("Consecutive-trade events: %d", len(events))
    return events


def _norm(s) -> str:
    return re.sub(r"[^a-z0-9]", "", str(s).lower())


def _find_col(columns, include, exclude=()):
    for key in include:
        for c in columns:
            if _norm(c) == key:
                return c
    for key in include:
        for c in columns:
            n = _norm(c)
            if key in n and not any(x in n for x in exclude):
                return c
    return None


def read_master_file(path: Path):
    try:
        if path.suffix.lower() == ".csv":
            raw = pd.read_csv(path, header=None, names=range(60), dtype=str, encoding_errors="replace")
        else:
            raw = pd.read_excel(path, header=None, dtype=str)
    except Exception as e:
        log.error("Could not read security master %s: %s", path.name, e)
        return None
    for i in range(min(len(raw), 40)):
        if raw.iloc[i].astype(str).str.strip().str.lower().eq("isin").any():
            df = raw.iloc[i + 1:].copy()
            df.columns = [str(c).strip() if pd.notna(c) else f"_blank{j}" for j, c in enumerate(raw.iloc[i])]
            return df.reset_index(drop=True)
    log.error("No header row containing 'ISIN' found in the first 40 rows of %s", path.name)
    return None


def parse_coupon(name: pd.Series, desc: pd.Series) -> pd.Series:
    pct = re.compile(r"(\d+(?:\.\d+)?)\s*%")

    def one(s):
        m = pct.search(s) if isinstance(s, str) else None
        return float(m.group(1)) if m else np.nan

    return name.map(one).fillna(desc.map(one))


def load_security_master(raw_dir: Path):
    files = sorted(p for p in raw_dir.rglob("*")
                   if p.is_file() and "master" in p.name.lower()
                   and p.suffix.lower() in (".csv", ".xlsx", ".xls") and not p.name.startswith("~$"))
    if not files:
        log.error("No security master file (a csv/xlsx with 'master' in its name) found under %s. "
                  "Duration, residual maturity and the government-adjusted shock will be empty.", raw_dir)
        return None
    log.info("Security master file(s): %s", [f.name for f in files])

    frames = [f for f in (read_master_file(p) for p in files) if f is not None]
    if not frames:
        return None
    raw = pd.concat(frames, ignore_index=True)
    cols = list(raw.columns)

    isin_c = _find_col(cols, ["isin"])
    mat_c = _find_col(cols, ["maturitydate", "dateofmaturity", "redemptiondate", "maturity"],
                      exclude=("residual", "remaining", "tenor"))
    if isin_c is None or mat_c is None:
        log.error("Security master needs an ISIN and a maturity-date column. Columns found: %s. "
                  "Duration and the government-adjusted shock will be empty.", cols)
        return None

    name_c = _find_col(cols, ["issuename"])
    desc_c = _find_col(cols, ["issuedesc", "issuedescription"])
    cpn_c = _find_col(cols, ["couponrate", "coupon"], exclude=("frequency", "date", "step"))
    freq_c = _find_col(cols, ["couponfrequency", "frequency"])
    type_c = _find_col(cols, ["sectype", "securitytype"])
    issuer_c = _find_col(cols, ["issuer"])
    issued_c = _find_col(cols, ["issuedate"])

    master = pd.DataFrame({"isin": raw[isin_c].astype(str).str.strip()})
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        master["maturity_date"] = pd.to_datetime(raw[mat_c], errors="coerce", dayfirst=True).astype("datetime64[ns]")
        master["issue_date"] = (pd.to_datetime(raw[issued_c], errors="coerce", dayfirst=True).astype("datetime64[ns]")
                                if issued_c else pd.NaT)

    if cpn_c is not None:
        coupon = pd.to_numeric(raw[cpn_c], errors="coerce")
    else:
        empty = pd.Series(np.nan, index=raw.index, dtype=object)
        coupon = parse_coupon(raw[name_c] if name_c else empty, raw[desc_c] if desc_c else empty)

    if freq_c is not None:
        freq_txt = raw[freq_c].astype(str).map(_norm)
        master["coupon_freq"] = freq_txt.map(FREQ_MAP)
        coupon = coupon.where(~((freq_txt == "0") & coupon.isna()), 0.0)
    else:
        master["coupon_freq"] = np.nan
    master["coupon"] = coupon
    master["sec_type"] = raw[type_c].astype(str).str.strip().str.upper() if type_c else np.nan
    master["issuer"] = raw[issuer_c].astype(str).str.strip() if issuer_c else np.nan

    master = master.drop_duplicates("isin", keep="last")
    live = master[~master["sec_type"].isin(EXCLUDED_SEC_TYPES)] if type_c else master
    log.info("Security master: %d ISINs (%d excl. CP/CD); maturity present %.1f%%, coupon present %.1f%%, "
             "coupon frequency present %.1f%%", len(master), len(live),
             100 * live["maturity_date"].notna().mean(), 100 * live["coupon"].notna().mean(),
             100 * live["coupon_freq"].notna().mean())
    return master


def modified_duration(y_pct, c_pct, n_years, freq):
    m = np.asarray(freq, dtype=float)
    yp = np.asarray(y_pct, dtype=float) / 100 / m
    cp = np.asarray(c_pct, dtype=float) / 100 / m
    n = np.asarray(n_years, dtype=float)
    periods = n * m
    with np.errstate(all="ignore"):
        mac_periods = (1 + yp) / yp - ((1 + yp) + periods * (cp - yp)) / (cp * ((1 + yp) ** periods - 1) + yp)
        dur = mac_periods / m / (1 + yp)
    dur[~np.isfinite(dur) | (dur <= 0) | (dur > n + 1)] = np.nan
    return dur


def fiscal_year_start(dates: pd.Series) -> pd.Series:
    return dates.dt.year - (dates.dt.month < 4).astype(int)


def attach_bond_terms(events: pd.DataFrame, master) -> pd.DataFrame:
    if master is None:
        for col in ["maturity_date", "issue_date", "coupon", "coupon_freq", "sec_type", "issuer",
                    "residual_maturity", "duration", "mat_fy"]:
            events[col] = np.nan
        return events
    events = events.merge(master, on="isin", how="left")
    events["residual_maturity"] = (events["maturity_date"] - events["date"]).dt.days / 365.25
    events.loc[events["residual_maturity"] <= 0, ["residual_maturity", "maturity_date"]] = np.nan
    coupon = events["coupon"].fillna(events["vwap_yield"])
    freq = events["coupon_freq"].fillna(1)
    events["duration"] = modified_duration(events["vwap_yield"], coupon, events["residual_maturity"], freq)
    events["mat_fy"] = fiscal_year_start(events["maturity_date"])
    log.info("Bond terms attached: maturity coverage %.1f%%, duration coverage %.1f%%",
             100 * events["maturity_date"].notna().mean(), 100 * events["duration"].notna().mean())
    return events


def _num(v):
    if isinstance(v, bool) or not isinstance(v, (int, float)) or v != v:
        return None
    return float(v)


def parse_rbi_gsec(path):
    if path is None:
        log.error("RBI G-sec workbook not found at %s. The government-adjusted shock will be empty.",
                  RBI_GSEC_XLSX)
        return None
    try:
        from openpyxl import load_workbook
    except ImportError:
        log.error("openpyxl is not installed. Run: pip install openpyxl. The government-adjusted "
                  "shock will be empty.")
        return None
    try:
        rows = list(load_workbook(path, read_only=True, data_only=True).active.iter_rows(values_only=True))
    except Exception as e:
        log.error("Could not read %s: %s", path.name, e)
        return None

    exact_re = re.compile(r"^\s*(\d{4})\s*-\s*\d{2,4}\s*$")
    above_re = re.compile(r"^\s*(\d{4})\s*&\s*above\s*$", re.I)
    columns, records = [], []

    for i, r in enumerate(rows):
        if len(r) < 4:
            continue
        if "State Government Securities" in [x for x in r if isinstance(x, str)]:
            end = r.index("State Government Securities")
            columns = []
            for label_row in rows[i + 1:i + 3]:
                found = []
                for pos in range(3, min(end, len(label_row))):
                    v = label_row[pos]
                    if isinstance(v, str):
                        m_exact, m_above = exact_re.match(v), above_re.match(v)
                        if m_exact:
                            found.append((pos, "exact", int(m_exact.group(1))))
                        elif m_above:
                            found.append((pos, "above", int(m_above.group(1))))
                if found:
                    columns = found
                    break
            continue

        if columns and isinstance(r[1], dt.datetime) and r[2] == "Amount" and i + 2 < len(rows):
            week_date = r[1]
            if week_date < GSEC_MIN_DATE:
                continue
            min_row, max_row = rows[i + 1], rows[i + 2]
            if min_row[2] != "YTM% Min" or max_row[2] != "YTM% Max":
                continue
            for pos, kind, year in columns:
                ymin, ymax = _num(min_row[pos]), _num(max_row[pos])
                if ymin is None or ymax is None or not (0 < ymin <= ymax < 25):
                    continue
                mid = (ymin + ymax) / 2
                if kind == "exact":
                    records.append((week_date, year, mid))
                else:
                    records.extend((week_date, fy, mid) for fy in range(year, year + 61))

    if not records:
        log.error("Parsed 0 usable rows from %s; the sheet layout may have changed. The "
                  "government-adjusted shock will be empty.", path.name)
        return None

    gsec = pd.DataFrame(records, columns=["g_date", "mat_fy", "gsec_yield"])
    gsec["g_date"] = pd.to_datetime(gsec["g_date"]).astype("datetime64[ns]")
    gsec["mat_fy"] = gsec["mat_fy"].astype("int64")
    gsec = gsec.drop_duplicates(["g_date", "mat_fy"], keep="last").sort_values("g_date").reset_index(drop=True)
    log.info("RBI G-sec: %d weekly bucket yields, %s to %s", len(gsec),
             gsec["g_date"].min().date(), gsec["g_date"].max().date())
    return gsec


def find_rbi_file():
    if RBI_GSEC_XLSX.exists():
        return RBI_GSEC_XLSX
    hits = [p for p in sorted(RAW_DIR.rglob("Secondary*Outright*.xlsx")) if not p.name.startswith("~$")]
    return hits[0] if hits else None


def lookup_gsec(dates, fys, gsec) -> np.ndarray:
    n = len(dates)
    out = np.full(n, np.nan)
    if gsec is None or n == 0:
        return out
    left = pd.DataFrame({
        "date": pd.DatetimeIndex(dates).astype("datetime64[ns]"),
        "mat_fy": np.asarray(fys, dtype=float),
        "pos": np.arange(n),
    }).dropna(subset=["date", "mat_fy"])
    if left.empty:
        return out
    left["mat_fy"] = left["mat_fy"].astype("int64")
    merged = pd.merge_asof(
        left.sort_values("date"), gsec, left_on="date", right_on="g_date", by="mat_fy",
        direction="forward", tolerance=pd.Timedelta(days=GSEC_MAX_FILL_DAYS),
    )
    out[merged["pos"].to_numpy()] = merged["gsec_yield"].to_numpy()
    return out


def attach_government_shock(events: pd.DataFrame, gsec) -> pd.DataFrame:
    events["g_t"] = lookup_gsec(events["date"], events["mat_fy"], gsec)
    events["g_t0"] = lookup_gsec(events["prev_date"], events["mat_fy"], gsec)
    events["gov_shock"] = events["raw_dy"] - (events["g_t"] - events["g_t0"])
    log.info("Government-adjusted shock coverage: %.1f%%", 100 * events["gov_shock"].notna().mean())
    return events


def attach_forward_returns(events: pd.DataFrame, panel: pd.DataFrame, gsec) -> pd.DataFrame:
    right = panel[["isin", "date", "vwap_yield"]].rename(columns={"date": "fwd_date", "vwap_yield": "fwd_yield"})
    right = right.sort_values("fwd_date")
    for h in HORIZONS:
        left = events[["isin", "date"]].copy()
        left["target"] = left["date"] + pd.Timedelta(days=h)
        left["pos"] = np.arange(len(left))
        merged = pd.merge_asof(left.sort_values("target"), right, left_on="target", right_on="fwd_date",
                                by="isin", direction="backward").sort_values("pos")
        valid = (merged["fwd_date"] > merged["date"]).to_numpy()
        fwd_date = merged["fwd_date"].where(valid)
        fwd_yield = merged["fwd_yield"].where(valid).to_numpy()

        events[f"fwd_days_{h}"] = (fwd_date - merged["date"]).dt.days.to_numpy()
        events[f"fwd_dy_{h}"] = fwd_yield - events["vwap_yield"].to_numpy()
        g_fwd = lookup_gsec(fwd_date.to_numpy(), events["mat_fy"], gsec)
        events[f"fwd_resid_dy_{h}"] = events[f"fwd_dy_{h}"] - (g_fwd - events["g_t"].to_numpy())
        events[f"er_{h}"] = -events["duration"] * events[f"fwd_resid_dy_{h}"]
    return events


def add_project_variables(events: pd.DataFrame) -> pd.DataFrame:
    events["year"] = events["date"].dt.year
    events["log1p_gap"] = np.log1p(events["gap"])
    events["abs_raw_dy"] = events["raw_dy"].abs()
    events["abs_shock"] = events["gov_shock"].abs()
    events["shock_positive"] = np.where(events["gov_shock"].isna(), np.nan, (events["gov_shock"] > 0).astype(float))
    for g in (20, 30, 60, 120):
        events[f"gap_ge_{g}"] = (events["gap"] >= g).astype(int)
    events["stale"] = (events["gap"] >= STALE_GAP).astype(int)
    events["gap_bucket"] = pd.cut(events["gap"], [0, 5, 10, 20, 30, 60, 120, 365, np.inf],
                                  labels=["1-5", "6-10", "11-20", "21-30", "31-60", "61-120", "121-365", "365+"])
    events["period"] = np.select(
        [events["date"] < "2020-03-01", events["date"] < "2021-01-01", events["date"] < "2023-01-01"],
        ["1_pre_covid", "2_crisis", "3_recovery"], default="4_post_recovery")
    events["duration_group"] = pd.cut(events["duration"], [0, 1.5, 4, np.inf], labels=["short", "medium", "long"])
    events["duration_quintile"] = pd.qcut(events["duration"], 5, labels=False, duplicates="drop") + 1
    return events


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    master = load_security_master(RAW_DIR)
    panel = drop_excluded_types(clean_panel(load_cbm_panel(RAW_DIR)), master)
    events = build_events(panel)
    gsec = parse_rbi_gsec(find_rbi_file())

    events = attach_bond_terms(events, master)
    events = attach_government_shock(events, gsec)
    events = attach_forward_returns(events, panel, gsec)
    events = add_project_variables(events)

    events.to_csv(OUT_DIR / "events_full.csv", index=False)
    h2 = events[(events["gap"] >= STALE_GAP) & events["gov_shock"].notna() & events["er_20"].notna()]
    h2.to_csv(OUT_DIR / "h2_sample.csv", index=False)

    log.info("Saved events_full.csv (%d rows) and h2_sample.csv (%d rows) to %s", len(events), len(h2), OUT_DIR)
    log.info("Coverage: duration %.1f%%, gov_shock %.1f%%, er_20 %.1f%%",
             100 * events["duration"].notna().mean(), 100 * events["gov_shock"].notna().mean(),
             100 * events["er_20"].notna().mean())


if __name__ == "__main__":
    main()